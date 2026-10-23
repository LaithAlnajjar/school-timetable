"""Launch an EXE with development paths removed, then inspect its produced artifacts."""

import argparse
import json
import os
import subprocess
from pathlib import Path

import fitz
import pefile
from openpyxl import load_workbook
from pypdf import PdfReader


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("executable", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    exe, out = args.executable.resolve(), args.output.resolve()
    pe = pefile.PE(str(exe), fast_load=True)
    assert pe.FILE_HEADER.Machine == 0x8664, "Expected Windows x64"
    assert pe.OPTIONAL_HEADER.Subsystem == 2, "Expected a GUI executable without a console"
    pe.close()
    assert not (exe.parent / "_internal" / "icuuc.dll").exists(), (
        "A foreign ICU DLL shadows Windows ICU"
    )
    env = dict(os.environ)
    for key in list(env):
        if key.startswith(("PYTHON", "QT_")) or key == "VIRTUAL_ENV":
            env.pop(key)
    windows_dir = Path(os.environ["SystemRoot"])
    env["PATH"] = os.pathsep.join([str(windows_dir / "System32"), str(windows_dir)])
    env["QT_QPA_PLATFORM"] = "windows"
    result = subprocess.run(
        [str(exe), "--smoke-test", str(out)],
        cwd=str(exe.parent),
        env=env,
        timeout=100,
        creationflags=subprocess.CREATE_NO_WINDOW,
        check=False,
    )
    report_path = out / "smoke-report.json"
    if not report_path.exists():
        raise RuntimeError(f"Packaged app exited {result.returncode} without a smoke report")
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert result.returncode == 0 and report["passed"], report
    assert report["frozen"] and report["qt_platform"] == "windows", report
    for kind, count in (("classes", 3), ("teachers", 5)):
        workbook = load_workbook(out / f"{kind}.xlsx")
        assert len(workbook.worksheets) == count + 1
        assert all(ws.max_row >= 5 for ws in workbook.worksheets[1:])
        pdf = out / f"{kind}.pdf"
        pages = PdfReader(pdf).pages
        assert len(pages) >= count
        text = " ".join(" ".join(page.extract_text().split()) for page in pages)
        assert "Willow Secondary" in text and "Free period" in text
        with fitz.open(pdf) as document:
            pix = document[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            assert len(set(pix.samples)) > 20, "Rendered PDF appears blank"
            pix.save(str(out / f"{kind}-page1.png"))
    print(
        json.dumps(
            {
                "executable": str(exe),
                "passed": True,
                "elapsed_seconds": report["elapsed_seconds"],
                "ui_heartbeat_events": report["ui_heartbeat_events"],
                "qt_platform": report["qt_platform"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
