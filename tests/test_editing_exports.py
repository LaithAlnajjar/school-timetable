from copy import deepcopy

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from school_timetable.editing import EditError, ScheduleHistory, move_or_swap
from school_timetable.exports import ExportError, export_excel, export_pdf, sheet_title
from school_timetable.models import (
    Day,
    Lesson,
    Project,
    Requirement,
    SchoolClass,
    Slot,
    Subject,
    Teacher,
)
from school_timetable.validation import validate_schedule
from school_timetable.views import timetable


def editable():
    return Project(
        school_name="Example school",
        days=[
            Day(
                "d",
                "Sunday",
                [
                    Slot("s1", "1"),
                    Slot("s2", "2", "08:45", "09:30"),
                    Slot("s3", "3", "09:30", "10:15"),
                ],
            )
        ],
        teachers=[Teacher("t1", "Alice"), Teacher("t2", "Bob")],
        classes=[SchoolClass("c", "7A")],
        subjects=[Subject("a", "Art"), Subject("b", "Biology")],
        requirements=[Requirement("r1", "c", "a", "t1"), Requirement("r2", "c", "b", "t2")],
        lessons=[Lesson("r1:single:0", "r1", ["s1"]), Lesson("r2:single:0", "r2", ["s2"])],
    )


def test_moves_swaps_atomic_rejection_and_undo_redo():
    p = editable()
    assert validate_schedule(p) == []
    old = deepcopy(p.lessons)
    moved = move_or_swap(p, old[0].id, "s3")
    assert p.lessons == old
    assert moved[0].slot_ids == ["s3"]
    history = ScheduleHistory()
    history.apply(p, moved)
    history.undo(p)
    assert p.lessons == old
    history.redo(p)
    assert p.lessons == moved
    swapped = move_or_swap(p, p.lessons[0].id, "s2", p.lessons[1].id)
    assert swapped[0].slot_ids == ["s2"]
    assert swapped[1].slot_ids == ["s3"]
    before = deepcopy(p)
    with pytest.raises(EditError, match="two lessons"):
        move_or_swap(p, p.lessons[0].id, "s2")
    assert p == before
    p.lessons[0].locked = True
    with pytest.raises(EditError, match="Unlock"):
        move_or_swap(p, p.lessons[0].id, "s1")


def test_double_moves_as_block_and_invalid_swap_rejected():
    p = editable()
    p.requirements[0].periods = 2
    p.requirements[0].doubles = 1
    p.lessons[0] = Lesson("r1:double:0", "r1", ["s1", "s2"])
    p.lessons[1].slot_ids = ["s3"]
    assert validate_schedule(p) == []
    # Moving the double to the last period cannot fit; nothing changes.
    before = deepcopy(p)
    with pytest.raises(EditError, match="does not fit"):
        move_or_swap(p, p.lessons[0].id, "s3")
    with pytest.raises(EditError):
        move_or_swap(p, p.lessons[0].id, "s3", p.lessons[1].id)
    assert p == before
    # Leave enough space between a single and a double for a complete atomic swap.
    p.days[0].slots.append(Slot("s4", "4", "10:15", "11:00"))
    p.lessons[0].slot_ids = ["s3", "s4"]
    p.lessons[1].slot_ids = ["s1"]
    swapped = move_or_swap(p, p.lessons[0].id, "s1", p.lessons[1].id)
    assert swapped[0].slot_ids == ["s1", "s2"]
    assert swapped[1].slot_ids == ["s3"]
    assert validate_schedule(p, swapped) == []


def test_class_teacher_views_and_current_edited_export(tmp_path, qapp):
    p = editable()
    p.lessons = move_or_swap(p, p.lessons[0].id, "s3")
    class_view = timetable(p, "classes", "c")
    teacher_view = timetable(p, "teachers", "t1")
    assert class_view.columns[0][2].lesson_id == teacher_view.columns[0][2].lesson_id
    assert "Art" in class_view.columns[0][2].text
    assert "Alice" in class_view.columns[0][2].text
    assert "7A" in teacher_view.columns[0][2].text
    xlsx = tmp_path / "classes.xlsx"
    pdf = tmp_path / "teachers.pdf"
    export_excel(p, xlsx)
    export_pdf(p, pdf, "teachers")
    wb = load_workbook(xlsx)
    assert "Free period" in wb["7A"]["A5"].value
    assert "Art" in wb["7A"]["A7"].value
    assert wb["7A"].print_title_rows == "$1:$4"
    pages = PdfReader(pdf).pages
    assert len(pages) == 2
    text = " ".join("\n".join(page.extract_text() for page in pages).split())
    for expected in ("Example school", "Alice", "Bob", "Sunday", "08:00", "Art", "Free period"):
        assert expected in text


def test_invalid_exports_create_no_files(tmp_path, qapp):
    p = editable()
    p.lessons.pop()
    for export, suffix in ((export_excel, "xlsx"), (export_pdf, "pdf")):
        target = tmp_path / f"invalid.{suffix}"
        with pytest.raises(ExportError, match="exactly"):
            export(p, target)
        assert not target.exists()


def test_sheet_names_and_literal_formula_names(tmp_path):
    used = {"index"}
    titles = [
        sheet_title(name, used) for name in ("Index", "A/B:*?[]", "a_b____", "x" * 80, "x" * 80)
    ]
    assert all(len(name) <= 31 for name in titles)
    assert len({name.casefold() for name in titles}) == len(titles)
    p = editable()
    p.school_name = "=1+1"
    p.classes[0].name = '=HYPERLINK("example")'
    target = tmp_path / "literal.xlsx"
    export_excel(p, target)
    wb = load_workbook(target)
    assert wb["Index"]["A1"].value == "=1+1"
    assert wb["Index"]["A1"].data_type == "s"


def test_long_names_and_many_rows_repeat_pdf_headers(tmp_path, qapp):
    p = editable()
    p.classes[0].name = "Long class name " * 10
    p.teachers[0].name = "Teacher with a long descriptive name " * 4
    # Repeat many empty timed days, preserving globally unique slot IDs.
    for i in range(1, 8):
        day = deepcopy(p.days[0])
        day.id, day.name = f"d{i}", f"School day {i}"
        for slot in day.slots:
            slot.id += str(i)
        p.days.append(day)
    target = tmp_path / "long.pdf"
    export_pdf(p, target)
    pages = PdfReader(target).pages
    assert len(pages) == 3  # 3 + 3 + 2 day columns, fixed readable font
    assert all("Example school" in " ".join(page.extract_text().split()) for page in pages)


def test_vertical_pdf_pagination_keeps_all_periods(tmp_path, qapp):
    p = editable()
    for i in range(20):
        start = 10 * 60 + 15 + i * 30
        end = start + 30
        p.days[0].slots.append(
            Slot(
                f"extra{i}",
                f"Extra period {i}",
                f"{start // 60:02}:{start % 60:02}",
                f"{end // 60:02}:{end % 60:02}",
            )
        )
    target = tmp_path / "tall.pdf"
    export_pdf(p, target)
    pages = PdfReader(target).pages
    assert len(pages) >= 2
    text = " ".join(" ".join(page.extract_text().split()) for page in pages)
    assert "Extra period 19" in text
    assert all("Example school" in " ".join(page.extract_text().split()) for page in pages)
